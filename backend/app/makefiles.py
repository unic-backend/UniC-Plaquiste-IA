"""Fichiers créés par l'IA à la demande du patron : un PDF (lettre, rapport, compte rendu, attestation…) ou un tableur Excel.

Le contenu vient du patron ou de l'IA ; ce module ne fait que le mettre en forme, sans rien ajouter. Le fichier devient un
document de la conversation (téléchargeable, partageable) et peut être joint à un brouillon d'e-mail. Rien ne part tout seul.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy.orm import Session

from app.config import settings
from app.models import Artifact, new_id
from app.services import company_dict, store_artifact

MAX_PDF_CHARS = 60_000
MAX_SHEETS, MAX_ROWS, MAX_COLS = 20, 5000, 60
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class MakeFileError(ValueError):
    pass


def _filename(title: str, ext: str) -> str:
    stem = re.sub(r"[^\w\- ]+", "", title, flags=re.UNICODE).strip().replace(" ", "_")[:60] or "document"
    return f"{stem}.{ext}"


def _inline(text: str) -> str:
    """Texte sûr pour reportlab : échappé, **gras** et *italique* gardés."""
    t = escape(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    return re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", t)


def _out(kind: str, filename: str) -> Path:
    folder = settings.artifacts_path / "made"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{Path(filename).stem}_{new_id()[:8]}.{kind}"


CHART_MARK = re.compile(r"^\[\[graphique (\d+)\]\]$")


def _bar_chart(spec: dict, width: float):
    """Histogramme groupé (reportlab) : {"titre", "labels": [...], "series": [{"nom", "valeurs": [...], "couleur": "#hex"}]}."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    labels = [str(x)[:12] for x in spec.get("labels") or []]
    series = [s for s in spec.get("series") or [] if len(s.get("valeurs") or []) == len(labels)]
    if not labels or not series:
        return None
    d = Drawing(width, 200)
    d.add(String(0, 188, str(spec.get("titre") or "")[:90], fontName="Helvetica-Bold", fontSize=10))
    bc = VerticalBarChart()
    bc.x, bc.y, bc.width, bc.height = 46, 38, width - 56, 132
    bc.data = [[float(v or 0) for v in s["valeurs"]] for s in series]
    bc.categoryAxis.categoryNames = labels
    bc.categoryAxis.labels.fontName = bc.valueAxis.labels.fontName = "Helvetica"
    bc.categoryAxis.labels.fontSize = 7
    bc.categoryAxis.labels.angle = 30 if len(labels) > 6 else 0
    bc.categoryAxis.labels.dy = -8 if len(labels) > 6 else -2
    bc.valueAxis.labels.fontSize = 7
    bc.valueAxis.valueMin = min(0.0, *(v for row in bc.data for v in row))
    bc.valueAxis.labelTextFormat = lambda v: f"{v / 1e6:.1f} M" if abs(v) >= 1e6 else f"{v / 1e3:.0f} k" if abs(v) >= 1e3 else f"{v:.0f}"
    bc.groupSpacing, bc.barSpacing = 6, 1
    for i, s in enumerate(series):
        bc.bars[i].fillColor = colors.HexColor(s.get("couleur") or ("#B86B2F", "#2E7D4F", "#6B7280")[i % 3])
        bc.bars[i].strokeColor = None
    d.add(bc)
    lg = Legend()
    lg.fontName, lg.fontSize, lg.columnMaximum, lg.dx, lg.dy, lg.deltax = "Helvetica", 8, 1, 8, 8, 72
    lg.x, lg.y = width - 72 * len(series), 192
    lg.colorNamePairs = [(bc.bars[i].fillColor, str(s.get("nom") or "")[:20]) for i, s in enumerate(series)]
    d.add(lg)
    return d


def make_pdf(db: Session, title: str, content: str, user_id: str | None, charts: list[dict] | None = None) -> Artifact:
    """PDF A4 : en-tête de l'entreprise, titre, date, texte (titres « # », listes « - », **gras**), pages numérotées.
    `charts` : histogrammes placés là où le texte porte « [[graphique N]] » (N à partir de 1), sinon à la fin."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    title, content = (title or "").strip(), (content or "").strip()
    if not title or not content:
        raise MakeFileError("Il faut un titre et un contenu.")
    if len(content) > MAX_PDF_CHARS:
        raise MakeFileError(f"Contenu trop long ({len(content)} caractères, {MAX_PDF_CHARS} au plus).")
    company = company_dict(db)
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["Normal"], fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=17, spaceBefore=8, spaceAfter=4)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    story: list = [Paragraph(_inline(title), h1), Paragraph(datetime.now().strftime("%d/%m/%Y"), small), Spacer(1, 6 * mm)]
    bullets: list = []

    def flush() -> None:
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(b, body), leftIndent=10) for b in bullets], bulletType="bullet", start="•"))
            bullets.clear()

    width = A4[0] - 36 * mm
    placed: set[int] = set()
    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        mark = CHART_MARK.match(line)
        if mark:
            flush()
            n = int(mark.group(1)) - 1
            if charts and 0 <= n < len(charts) and (drawing := _bar_chart(charts[n], width)) is not None:
                story.append(drawing)
                placed.add(n)
            continue
        if re.match(r"^[-*•]\s+", line):
            bullets.append(_inline(re.sub(r"^[-*•]\s+", "", line)))
            continue
        flush()
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        story.append(Paragraph(_inline(m.group(2)), h2) if m else Paragraph(_inline(line), body))
    flush()
    for n, spec in enumerate(charts or []):
        if n not in placed and (drawing := _bar_chart(spec, width)) is not None:
            story.append(drawing)

    head = " · ".join(x for x in (company.get("name"), company.get("phone"), company.get("email")) if x)

    def decorate(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawString(18 * mm, A4[1] - 12 * mm, head[:110])
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    name = _filename(title, "pdf")
    dest = _out("pdf", name)
    SimpleDocTemplate(str(dest), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm, bottomMargin=16 * mm,
                      title=title, author=company.get("name") or "").build(story, onFirstPage=decorate, onLaterPages=decorate)
    art = store_artifact(db, dest, name, "made_file", dest.stem, f"made-{dest.stem}", user_id)
    db.commit()
    return art


def _cell(v):
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v)
    return s[:32000]


def _excel_chart(worksheets: list, spec: dict) -> None:
    """Graphique natif (modifiable dans Excel) à droite du tableau ; une demande incohérente est ignorée, jamais une erreur."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    try:
        ws = worksheets[int(spec.get("feuille") or 1) - 1]
        x, ys = int(spec.get("x") or 1), [int(c) for c in (spec.get("y") or [2])]
    except (IndexError, TypeError, ValueError):
        return
    if ws.max_row < 2 or x > ws.max_column or not all(1 <= c <= ws.max_column for c in ys):
        return
    chart = LineChart() if spec.get("type") == "courbe" else BarChart()
    chart.title = str(spec.get("titre") or "")[:80] or None
    chart.height, chart.width = 9, 18
    for c in ys:
        chart.add_data(Reference(ws, min_col=c, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=x, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, f"{get_column_letter(ws.max_column + 2)}2")


def make_spreadsheet(db: Session, title: str, sheets: list, user_id: str | None, charts: list[dict] | None = None) -> Artifact:
    """Classeur Excel : une feuille par entrée {name, rows} ; la 1re ligne est l'en-tête (gras, figée) ; « =… » reste une formule.
    `charts` : graphiques Excel natifs {"feuille": 1, "type": "barres"|"courbe", "titre", "x": 1, "y": [2]} (colonnes à partir de 1)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    title = (title or "").strip()
    if not title:
        raise MakeFileError("Il faut un titre.")
    if not isinstance(sheets, list) or not sheets:
        raise MakeFileError("Il faut au moins une feuille avec des lignes.")
    if len(sheets) > MAX_SHEETS:
        raise MakeFileError(f"Trop de feuilles ({MAX_SHEETS} au plus).")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for i, sh in enumerate(sheets, start=1):
        if not isinstance(sh, dict) or not isinstance(sh.get("rows"), list) or not sh["rows"]:
            raise MakeFileError(f"Feuille {i} : il faut une liste de lignes.")
        rows = sh["rows"]
        if len(rows) > MAX_ROWS or any(not isinstance(r, list) or len(r) > MAX_COLS for r in rows):
            raise MakeFileError(f"Feuille {i} : {MAX_ROWS} lignes et {MAX_COLS} colonnes au plus, chaque ligne étant une liste.")
        name = re.sub(r"[\[\]:*?/\\]", " ", str(sh.get("name") or f"Feuille {i}")).strip()[:31] or f"Feuille {i}"
        while name in used:
            name = f"{name[:28]} {i}"
        used.add(name)
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([_cell(v) for v in r])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in range(1, max(len(r) for r in rows) + 1):
            width = max((len(str(r[col - 1])) for r in rows if len(r) >= col and r[col - 1] is not None), default=8)
            ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)
    for spec in charts or []:
        _excel_chart(wb.worksheets, spec)
    name = _filename(title, "xlsx")
    dest = _out("xlsx", name)
    wb.save(str(dest))
    art = store_artifact(db, dest, name, "made_file", dest.stem, f"made-{dest.stem}", user_id, mime=XLSX_MIME)
    db.commit()
    return art
