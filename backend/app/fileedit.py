"""Modifier un fichier reçu (PDF, Word, Excel) : remplacer, supprimer ou ajouter du texte, puis VÉRIFIER le résultat.

Le fichier d'origine n'est jamais touché : une copie « (modifié) » est créée et proposée au patron (téléchargement, partage).
PDF : le texte remplacé est réellement retiré du fichier (pas de rectangle blanc posé par-dessus : l'ancien texte ne reste pas copiable).
Limites dites franchement : une mise en page ne se recompose pas (pas de ligne de tableau ajoutée/supprimée dans un PDF, les traits restent) ;
un PDF scanné (image) n'a pas de texte à modifier ; les totaux ne se recalculent pas tout seuls : l'assistant les corrige lui-même.
"""
from __future__ import annotations

import ctypes
import re
from pathlib import Path

from sqlalchemy.orm import Session

from app.config import settings
from app.models import StoredFile
from app.services import store_artifact

MAX_BLOCKS = 500


class FileEditError(Exception):
    pass


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace(" ", " ").replace("\xa0", " ")).strip()


def _kind(rec: StoredFile) -> str:
    name = (rec.filename or "").lower()
    if name.endswith(".pdf") or rec.mime_type == "application/pdf":
        return "pdf"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    return ""


def _load(db: Session, file_id: str) -> tuple[StoredFile, Path, str]:
    rec = db.get(StoredFile, file_id) if file_id else None
    if rec is None:
        raise FileEditError("Fichier introuvable : aucun fichier joint à cette conversation.")
    path = Path(rec.path)
    if not path.exists():
        raise FileEditError("Le fichier n'est plus sur le serveur : demande de le renvoyer.")
    kind = _kind(rec)
    if not kind:
        raise FileEditError(f"Format non modifiable ({rec.filename}). Je sais modifier les PDF, les Word (.docx) et les Excel (.xlsx).")
    return rec, path, kind


# ---------------------------------------------------------------- PDF (pdfium : objets texte réellement retirés)

def _pdf_blocks(path: Path) -> list[dict]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    pdf = pdfium.PdfDocument(str(path))
    out: list[dict] = []
    for pn in range(len(pdf)):
        page = pdf[pn]
        tp = page.get_textpage()
        for i, o in enumerate(page.get_objects(filter=[raw.FPDF_PAGEOBJ_TEXT])):
            o.textpage = tp
            t = _norm(o.extract())
            if t:
                out.append({"page": pn + 1, "i": i, "text": t})
    return out


def _font_bytes(bold: bool) -> bytes:
    import os

    import reportlab
    return (Path(os.path.dirname(reportlab.__file__)) / "fonts" / ("VeraBd.ttf" if bold else "Vera.ttf")).read_bytes()


_NUMERIC = re.compile(r"^[\d\s.,]+(?:FCFA|F CFA|CFA|€|%|m²|m2)?$")


def _pdf_edit(src: Path, dst: Path, edits: list[dict]) -> list[dict]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    pdf = pdfium.PdfDocument(str(src))
    keep: list = []   # les tampons de police doivent vivre jusqu'à l'enregistrement
    fonts: dict[bool, int] = {}

    def font(bold: bool) -> int:
        if bold not in fonts:
            data = _font_bytes(bold)
            buf = (ctypes.c_uint8 * len(data)).from_buffer_copy(data)
            keep.append(buf)
            fonts[bold] = raw.FPDFText_LoadFont(pdf.raw, buf, len(data), raw.FPDF_FONT_TRUETYPE, True)
        return fonts[bold]

    def new_text(page, text: str, size: float, bold: bool, x: float, y: float, color=(0, 0, 0), right_edge: float | None = None):
        obj = raw.FPDFPageObj_CreateTextObj(pdf.raw, font(bold), size)
        codes = [ord(c) if ord(c) < 0x10000 else 63 for c in text]
        arr = (ctypes.c_ushort * (len(codes) + 1))(*codes, 0)
        raw.FPDFText_SetText(obj, arr)
        raw.FPDFPageObj_SetFillColor(obj, *color, 255)
        if right_edge is not None:   # montant aligné à droite comme l'ancien
            l, b, r, t = (ctypes.c_float() for _ in range(4))
            raw.FPDFPageObj_GetBounds(obj, l, b, r, t)
            x = right_edge - (r.value - l.value)
        raw.FPDFPageObj_Transform(obj, 1, 0, 0, 1, x, y)
        raw.FPDFPage_InsertObject(page.raw, obj)

    report: list[dict] = []
    for n, ed in enumerate(edits):
        op = ed.get("op", "replace")
        if op == "add_text":
            pn = max(1, int(ed.get("page") or 1)) - 1
            if pn >= len(pdf):
                report.append({"n": n, "ok": False, "why": "Page inexistante."})
                continue
            page = pdf[pn]
            lines = [l for l in str(ed.get("text") or "").split("\n") if l.strip()]
            size = float(ed.get("size") or 9)
            x = float(ed.get("x") or 42)
            y0 = float(ed.get("y") or 36)
            for k, line in enumerate(lines):
                new_text(page, line, size, False, x, y0 + (len(lines) - 1 - k) * size * 1.35)
            page.gen_content()
            report.append({"n": n, "ok": True, "count": len(lines), "op": op})
            continue
        find, repl = _norm(str(ed.get("find") or "")), str(ed.get("replace") if op == "replace" else "")
        if not find:
            report.append({"n": n, "ok": False, "why": "Texte à trouver manquant."})
            continue
        every = bool(ed.get("all", True))
        done = 0
        for pn in range(len(pdf)):
            page = pdf[pn]
            tp = page.get_textpage()
            objs = list(page.get_objects(filter=[raw.FPDF_PAGEOBJ_TEXT]))
            changed = False
            for o in objs:
                if done and not every:
                    break
                o.textpage = tp
                txt = _norm(o.extract())
                if find not in txt:
                    continue
                new = _norm(txt.replace(find, repl))
                l, b, r, t = o.get_bounds()
                mat = o.get_matrix()
                size = o.get_font_size() or 9
                try:
                    bold = (o.get_font().get_weight() or 400) >= 700 or "bold" in (o.get_font().get_base_name() or "").lower()
                except Exception:
                    bold = False
                cr, cg, cb, ca = (ctypes.c_uint() for _ in range(4))
                raw.FPDFPageObj_GetFillColor(o.raw, cr, cg, cb, ca)
                color = (cr.value, cg.value, cb.value)
                page.remove_obj(o)
                if new:
                    right = r if (_NUMERIC.match(txt) and _NUMERIC.match(new)) else None
                    new_text(page, new, size, bold, mat.e, mat.f, color, right)
                done += 1
                changed = True
            if changed:
                page.gen_content()
        report.append({"n": n, "ok": done > 0, "count": done, "op": op, "why": "" if done else "Texte introuvable dans un seul bloc : relis les blocs avec inspect_file et reprends le texte exact d'un bloc."})
    pdf.save(str(dst))
    return report


# ---------------------------------------------------------------- Word

def _docx_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def _docx_edit(src: Path, dst: Path, edits: list[dict]) -> list[dict]:
    import docx
    doc = docx.Document(str(src))
    report: list[dict] = []
    for n, ed in enumerate(edits):
        op = ed.get("op", "replace")
        if op == "add_paragraph":
            text = str(ed.get("text") or "")
            after = _norm(str(ed.get("after") or ""))
            target = next((p for p in _docx_paragraphs(doc) if after and after in _norm(p.text)), None) if after else None
            if target is not None:
                from copy import deepcopy
                el = deepcopy(target._p)
                target._p.addnext(el)
                from docx.text.paragraph import Paragraph
                para = Paragraph(el, target._parent)
                for r in para.runs[1:]:
                    r._r.getparent().remove(r._r)
                if para.runs:
                    para.runs[0].text = text
                else:
                    para.add_run(text)
            else:
                doc.add_paragraph(text)
            report.append({"n": n, "ok": True, "count": 1, "op": op})
            continue
        find, repl = str(ed.get("find") or ""), str(ed.get("replace") if op == "replace" else "")
        if not find:
            report.append({"n": n, "ok": False, "why": "Texte à trouver manquant."})
            continue
        every = bool(ed.get("all", True))
        done = 0
        for p in list(_docx_paragraphs(doc)):
            if done and not every:
                break
            if find not in p.text:
                continue
            hit = False
            for r in p.runs:
                if find in r.text:
                    r.text = r.text.replace(find, repl)
                    hit = True
            if not hit and p.runs:   # le texte est coupé entre plusieurs « runs » : on le recompose dans le premier
                full = p.text.replace(find, repl)
                p.runs[0].text = full
                for r in p.runs[1:]:
                    r.text = ""
            if op == "delete_paragraph" and not p.text.strip():
                p._p.getparent().remove(p._p)
            done += 1
        report.append({"n": n, "ok": done > 0, "count": done, "op": op, "why": "" if done else "Texte introuvable."})
    doc.save(str(dst))
    return report


# ---------------------------------------------------------------- Excel

def _xlsx_edit(src: Path, dst: Path, edits: list[dict]) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(src))   # les formules restent des formules (Excel recalcule à l'ouverture)
    report: list[dict] = []
    for n, ed in enumerate(edits):
        op = ed.get("op", "replace")
        try:
            ws = wb[ed["sheet"]] if ed.get("sheet") else wb.active
            if op == "set_cell":
                ws[str(ed["cell"])] = ed.get("value")
                report.append({"n": n, "ok": True, "count": 1, "op": op})
            elif op == "add_row":
                ws.append(list(ed.get("values") or []))
                report.append({"n": n, "ok": True, "count": 1, "op": op})
            else:
                find, repl = str(ed.get("find") or ""), str(ed.get("replace") if op == "replace" else "")
                done = 0
                for w in ([ws] if ed.get("sheet") else wb.worksheets):
                    for row in w.iter_rows():
                        for c in row:
                            if isinstance(c.value, str) and find and find in c.value and not c.value.startswith("="):
                                c.value = c.value.replace(find, repl)
                                done += 1
                report.append({"n": n, "ok": done > 0, "count": done, "op": op, "why": "" if done else "Texte introuvable dans une cellule."})
        except KeyError:
            report.append({"n": n, "ok": False, "why": "Feuille introuvable."})
        except Exception as exc:
            report.append({"n": n, "ok": False, "why": f"Modification impossible : {str(exc)[:120]}"})
    wb.save(str(dst))
    return report


def _text_of(path: Path, kind: str) -> str:
    if kind == "pdf":
        return _norm(" ".join(b["text"] for b in _pdf_blocks(path)))
    if kind == "docx":
        import docx
        d = docx.Document(str(path))
        return _norm(" ".join(p.text for p in _docx_paragraphs(d)))
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    return _norm(" ".join(str(c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if c.value is not None))


# ---------------------------------------------------------------- API de l'assistant

def inspect(db: Session, file_id: str) -> dict:
    rec, path, kind = _load(db, file_id)
    if kind == "pdf":
        blocks = _pdf_blocks(path)
        if not blocks:
            return {"fichier": rec.filename, "type": "pdf", "texte_modifiable": False,
                    "note": "PDF scanné ou image : aucun texte à modifier dans le fichier. Propose de REFAIRE le document (create_quote avec les lignes lues) au lieu de le modifier."}
        return {"fichier": rec.filename, "type": "pdf", "texte_modifiable": True, "blocs": blocks[:MAX_BLOCKS], "total_blocs": len(blocks),
                "note": "Chaque bloc est un morceau de texte (souvent une cellule). Pour remplacer, donne le texte exact du bloc ou une partie de ce bloc. "
                        "Les totaux ne se recalculent pas : modifie aussi chaque total qui dépend du changement."}
    if kind == "docx":
        import docx
        d = docx.Document(str(path))
        paras = [{"i": i, "text": _norm(p.text)} for i, p in enumerate(_docx_paragraphs(d)) if _norm(p.text)]
        return {"fichier": rec.filename, "type": "docx", "texte_modifiable": True, "paragraphes": paras[:MAX_BLOCKS], "total": len(paras)}
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    cells = [{"cell": f"{w.title}!{c.coordinate}", "value": str(c.value)[:120]} for w in wb.worksheets for row in w.iter_rows() for c in row if c.value is not None]
    return {"fichier": rec.filename, "type": "xlsx", "texte_modifiable": True, "feuilles": [w.title for w in wb.worksheets], "cellules": cells[:MAX_BLOCKS], "total": len(cells),
            "note": "Les formules (=…) restent des formules : Excel les recalcule à l'ouverture."}


def edit(db: Session, file_id: str, edits: list[dict], user_id: str | None) -> dict:
    rec, path, kind = _load(db, file_id)
    if not edits or not isinstance(edits, list):
        raise FileEditError("Aucune modification demandée.")
    if len(edits) > 60:
        raise FileEditError("Trop de modifications d'un coup (60 au plus).")
    if kind == "pdf" and not _pdf_blocks(path):
        raise FileEditError("PDF scanné ou image : aucun texte à modifier. Il faut refaire le document (create_quote).")
    folder = settings.artifacts_path / "edited"
    folder.mkdir(parents=True, exist_ok=True)
    from app.models import new_id
    key = new_id()[:8]
    stem = re.sub(r"[^\w\-]+", "_", Path(rec.filename).stem).strip("_")[:50] or "document"
    out = folder / f"{stem}_modifie_{key}{path.suffix.lower() or '.' + kind}"
    try:
        report = {"pdf": _pdf_edit, "docx": _docx_edit, "xlsx": _xlsx_edit}[kind](path, out, edits)
    except FileEditError:
        raise
    except Exception as exc:
        out.unlink(missing_ok=True)
        raise FileEditError(f"Modification impossible sur ce fichier ({type(exc).__name__}). Le fichier d'origine est intact.") from exc
    applied = [r for r in report if r.get("ok")]
    if not applied:
        out.unlink(missing_ok=True)
        return {"ok": False, "rapport": report, "note": "Rien n'a été modifié : aucun texte trouvé. Relis le fichier avec inspect_file et reprends le texte exact."}
    # vérification : le nouveau texte est bien dans la copie, l'ancien n'y est plus
    text = _text_of(out, kind)
    checks = []
    for ed, r in zip(edits, report):
        if not r.get("ok") or ed.get("op", "replace") in ("set_cell", "add_row"):
            continue
        op = ed.get("op", "replace")
        new = _norm(str(ed.get("replace") or ed.get("text") or "")) if op in ("replace", "add_text", "add_paragraph") else ""
        old = _norm(str(ed.get("find") or "")) if op in ("replace", "delete_text", "delete_paragraph") else ""
        if old and new and old in new:   # le nouveau texte contient l'ancien (« Faux plafond » -> « Faux plafond hydrofuge ») : rien à vérifier côté disparition
            old = ""
        checks.append({"n": r["n"], "nouveau_present": (new in text) if new else None,
                       "ancien_encore_present": (old in text) if old else None})
    mime = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[kind]
    art = store_artifact(db, out, f"{Path(rec.filename).stem} (modifié){path.suffix.lower()}", "edited_file", rec.id, f"edited-{key}", user_id, mime=mime)
    db.commit()
    return {"ok": True, "rapport": report, "verification": checks, "artifact": {"id": art.id, "filename": art.filename, "mime": mime, "size": art.size},
            "note": "Copie « (modifié) » créée, le fichier d'origine est intact. Relis la verification : si un nouveau texte est absent ou un ancien encore présent, corrige avant de répondre."}
