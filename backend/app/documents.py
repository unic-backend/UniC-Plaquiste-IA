"""Pipeline documents : extraction, classification, indexation. Pas d'envoi massif du PDF au modèle."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from pypdf import PdfReader
from sqlalchemy.orm import Session

from app import cad, ocr, pdfjob, vision
from app.config import settings
from app.models import DocumentChunk, ExtractedPage, StoredFile, new_id

DIM_RE = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*(?:m|mm|cm)\b|"
    r"(\d+(?:[.,]\d+)?)\s*[x×]\s*(\d+(?:[.,]\d+)?)\s*(?:m|mm)?",
    re.I,
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


class UploadRejected(ValueError):
    pass


ALLOWED_EXT = {".pdf", ".txt", ".csv", ".md", ".docx", ".xlsx", ".xlsm", ".dxf", ".ifc", ".dwg",
               ".jpg", ".jpeg", ".png", ".webp", ".gif"}
_TEXT_EXT = {".txt", ".csv", ".md", ".dxf", ".ifc"}
_SIGNATURES: dict[str, tuple[bytes, ...]] = {
    ".pdf": (b"%PDF",), ".docx": (b"PK\x03\x04",), ".xlsx": (b"PK\x03\x04",), ".xlsm": (b"PK\x03\x04",),
    ".jpg": (b"\xff\xd8\xff",), ".jpeg": (b"\xff\xd8\xff",), ".png": (b"\x89PNG\r\n\x1a\n",),
    ".gif": (b"GIF87a", b"GIF89a"), ".dwg": (b"AC10",),
}
_EXECUTABLE_HEADS = (b"MZ", b"\x7fELF", b"#!", b"\xca\xfe\xba\xbe", b"\xcf\xfa\xed\xfe")


def validate_upload(data: bytes, filename: str) -> str:
    """Contrôle réel d'un envoi : extension autorisée + contenu cohérent (signature) + pas d'exécutable. Rend l'extension."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXT:
        raise UploadRejected(f"Type de fichier non autorisé ({ext or 'sans extension'}).")
    if not data:
        raise UploadRejected("Fichier vide.")
    head = data[:16]
    if head.startswith(_EXECUTABLE_HEADS):
        raise UploadRejected("Fichier exécutable refusé.")
    if ext in _SIGNATURES and not head.startswith(_SIGNATURES[ext]):
        raise UploadRejected(f"Le contenu ne correspond pas à l'extension {ext}.")
    if ext == ".webp" and not (head[:4] == b"RIFF" and head[8:12] == b"WEBP"):
        raise UploadRejected("Le contenu ne correspond pas à l'extension .webp.")
    if ext in _TEXT_EXT and b"\x00" in data[:4096]:
        raise UploadRejected(f"Le contenu ne ressemble pas à du texte ({ext}).")
    return ext


def save_upload(data: bytes, filename: str, mime: str, user_id: str | None, project_id: str | None, db: Session) -> StoredFile:
    fid = new_id()
    filename = Path(filename or "fichier").name[:200] or "fichier"
    ext = validate_upload(data, filename)
    dest_dir = settings.uploads_path / fid[:2]
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{fid}{ext}"
    dest.write_bytes(data)
    rec = StoredFile(
        id=fid,
        filename=filename,
        mime_type=mime or "application/octet-stream",
        path=str(dest),
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        kind="upload",
        project_id=project_id,
        uploaded_by=user_id,
        processing_status="queued",
    )
    db.add(rec)
    db.commit()
    db.refresh(rec)
    return rec


def classify_page(text: str) -> str:
    t = (text or "").lower()
    if re.search(r"1\s*:\s*\d{2,4}|échelle|scale", t) or re.search(r"\b(plan|coupe|façade|elevation|niveau)\b", t):
        return "plan"
    if re.search(r"\b(devis|quotation|quote)\b", t):
        return "quotation"
    if re.search(r"\b(facture|invoice)\b", t):
        return "invoice"
    if re.search(r"\b(porte|door|fenêtre|window|cloison|partition)\b", t) and len(DIM_RE.findall(t)) >= 3:
        return "schedule"
    if len(t.strip()) < 40:
        return "scanned_or_empty"
    return "document"


MAX_OCR_PAGES = 15


OCR_BUDGET_S = 40   # temps total d'OCR par fichier : un gros scan ne bloque jamais l'envoi


def extract_pdf(file_rec: StoredFile, db: Session, max_pages: int = 2000) -> dict:
    import time

    path = Path(file_rec.path)
    heavy = ""
    try:   # texte lu dans un processus séparé à mémoire bornée (un plan lourd ne fait plus tomber le serveur)
        res = pdfjob.run("text", timeout=60, path=str(path), max_pages=max_pages)
        total, items = res["total"], res["pages"]
    except pdfjob.PdfJobError as exc:
        heavy = str(exc)
        try:
            total = len(PdfReader(str(path)).pages)
        except Exception:
            total = 0
        items = [{"t": "", "w": None, "h": None} for _ in range(min(total, max_pages))]
    n = len(items)
    file_rec.page_count = total
    db.query(ExtractedPage).filter(ExtractedPage.file_id == file_rec.id).delete()
    db.query(DocumentChunk).filter(DocumentChunk.file_id == file_rec.id).delete()

    pages_out = []
    empty_pages = 0
    started = time.monotonic()
    for i, item in enumerate(items):
        text = (item.get("t") or "").strip()
        if len(text) < 20 and not heavy and i < MAX_OCR_PAGES and time.monotonic() - started < OCR_BUDGET_S:
            text = ocr.ocr_pdf_page(path, i) or text   # la lecture visuelle (Claude) se fait à la demande : outil read_plan
        if len(text) < 20:
            empty_pages += 1
        klass = classify_page(text)
        db.add(ExtractedPage(file_id=file_rec.id, page_number=i + 1, text=text, classification=klass,
                             width=item.get("w"), height=item.get("h")))
        pages_out.append({"page": i + 1, "classification": klass, "chars": len(text)})
        for chunk in chunk_text(text, 900):
            db.add(DocumentChunk(file_id=file_rec.id, page_number=i + 1, text=chunk))

    ocr_needed = file_rec.page_count and empty_pages / max(n, 1) > 0.6
    file_rec.processing_status = "completed"
    if heavy:
        file_rec.processing_error = heavy + " Envoie une capture d'écran ou une photo du plan : je la lirai."
    if ocr_needed:
        file_rec.processing_status = "completed_no_ocr"
        file_rec.processing_error = file_rec.processing_error if heavy else (
            "La majorité des pages n'ont pas de texte lisible (plan dessiné ou scan). "
            + ("Demande « lis le plan » : je le regarde page par page." if vision.disponible()
               else "OCR et vision IA NON DISPONIBLES. Fournissez un PDF avec texte.")
        )
    db.commit()
    return {
        "file_id": file_rec.id,
        "pages": file_rec.page_count,
        "processed": n,
        "empty_pages": empty_pages,
        "ocr_needed": ocr_needed,
        "status": file_rec.processing_status,
        "warning": file_rec.processing_error,
        "page_summaries": pages_out[:50],
    }


def chunk_text(text: str, size: int = 900) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    parts = []
    buf = []
    n = 0
    for para in re.split(r"\n{2,}", text):
        para = para.strip()
        if not para:
            continue
        if n + len(para) > size and buf:
            parts.append("\n".join(buf))
            buf, n = [para], len(para)
        else:
            buf.append(para)
            n += len(para)
    if buf:
        parts.append("\n".join(buf))
    return parts


def search_pages(db: Session, file_id: str, query: str, limit: int = 20) -> list[dict]:
    q = (query or "").strip().lower()
    if not q:
        return []
    tokens = [t for t in re.split(r"\s+", q) if len(t) > 1]
    pages = (
        db.query(ExtractedPage)
        .filter(ExtractedPage.file_id == file_id)
        .order_by(ExtractedPage.page_number)
        .all()
    )
    scored = []
    for p in pages:
        blob = (p.text or "").lower()
        score = sum(blob.count(t) for t in tokens)
        if score:
            snippet = _snippet(p.text or "", tokens)
            scored.append({
                "page": p.page_number,
                "classification": p.classification,
                "score": score,
                "snippet": snippet,
            })
    scored.sort(key=lambda x: x["score"], reverse=True)
    return scored[:limit]


def extract_dimensions(text: str) -> list[str]:
    found = []
    for m in DIM_RE.finditer(text or ""):
        found.append(m.group(0))
    return found[:80]


def find_in_document(db: Session, file_id: str, topic: str) -> dict:
    topic_l = topic.lower()
    keywords_map = {
        "portes": ["porte", "door", "blk", "pe"],
        "doors": ["porte", "door"],
        "cloisons": ["cloison", "partition", "placo", "ba13", "doublage"],
        "partitions": ["cloison", "partition", "drywall"],
        "dimensions": ["dimension", "cote", "échelle", "1:", "mm", "m "],
        "quantités": ["quantit", "qté", "qty", "nbre", "nombre"],
        "quantities": ["quantit", "qty"],
    }
    keys = keywords_map.get(topic_l, [topic_l])
    pages = (
        db.query(ExtractedPage)
        .filter(ExtractedPage.file_id == file_id)
        .order_by(ExtractedPage.page_number)
        .all()
    )
    hits = []
    dims = []
    for p in pages:
        blob = (p.text or "").lower()
        if any(k in blob for k in keys):
            hits.append({
                "page": p.page_number,
                "classification": p.classification,
                "snippet": _snippet(p.text or "", keys),
                "dimensions": extract_dimensions(p.text or "")[:12],
            })
        dims.extend({"page": p.page_number, "value": d} for d in extract_dimensions(p.text or ""))
    return {
        "topic": topic,
        "hits": hits[:40],
        "hit_count": len(hits),
        "pages_with_hits": [h["page"] for h in hits],
        "dimensions_sample": dims[:40],
        "status": "ok" if hits else "none",
    }


def _snippet(text: str, tokens: list[str], radius: int = 140) -> str:
    low = text.lower()
    pos = -1
    for t in tokens:
        pos = low.find(t)
        if pos >= 0:
            break
    if pos < 0:
        return text[: radius * 2].replace("\n", " ")
    a = max(0, pos - radius)
    b = min(len(text), pos + radius)
    s = text[a:b].replace("\n", " ")
    if a > 0:
        s = "…" + s
    if b < len(text):
        s = s + "…"
    return s


def read_plaintext(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for name in wb.sheetnames[:12]:
        ws = wb[name]
        lines.append(f"# Feuille {name}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 200:
                lines.append("… (tronqué)")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                lines.append(" | ".join(vals))
    return "\n".join(lines)


def process_file(file_rec: StoredFile, db: Session) -> dict:
    # Déjà lu (à l'envoi du fichier) : on ne relit pas au moment du message. Évite doubles coûts OCR/vision et conflits de pages.
    if file_rec.processing_status in ("completed", "completed_no_ocr") and \
            db.query(ExtractedPage).filter(ExtractedPage.file_id == file_rec.id).first() is not None:
        return {"file_id": file_rec.id, "status": file_rec.processing_status, "pages": file_rec.page_count,
                "warning": file_rec.processing_error or None, "cached": True}
    db.query(ExtractedPage).filter(ExtractedPage.file_id == file_rec.id).delete()
    db.query(DocumentChunk).filter(DocumentChunk.file_id == file_rec.id).delete()
    path = Path(file_rec.path)
    mime = (file_rec.mime_type or "").lower()
    ext = path.suffix.lower()
    try:
        if ext == ".pdf" or "pdf" in mime:
            return extract_pdf(file_rec, db)
        text = ""
        if ext in {".txt", ".csv", ".md"}:
            text = read_plaintext(path)
        elif ext in {".docx"}:
            text = read_docx(path)
        elif ext in {".xlsx", ".xlsm"}:
            text = read_xlsx(path)
        elif ext in cad.CAD_EXT:
            text = cad.read_dxf(path) if ext == ".dxf" else cad.read_ifc(path)
        elif ext in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
            file_rec.processing_status = "completed"
            file_rec.page_count = 1
            read = ocr.ocr_image(path)
            seen = vision.describe_image(path)
            if seen:
                read = f"[Lecture visuelle IA] {seen}" + (f"\n\n[OCR]\n{read}" if read else "")
            db.add(ExtractedPage(
                file_id=file_rec.id, page_number=1,
                text=read or "[Image] Aucun texte lu. Analyse visuelle NON DISPONIBLE sans IA vision.",
                classification=classify_page(read) if read else "photo",
            ))
            for chunk in chunk_text(read):
                db.add(DocumentChunk(file_id=file_rec.id, page_number=1, text=chunk))
            db.commit()
            return {
                "file_id": file_rec.id,
                "kind": "image",
                "status": "completed",
                "chars": len(read),
                "warning": None if read else (
                    "Image stockée. Aucun texte lu : "
                    + ("pas de texte détecté." if ocr.disponible() or vision.disponible() else "OCR et vision IA NON DISPONIBLES (clé Claude absente).")
                ),
            }
        else:
            file_rec.processing_status = "unsupported"
            file_rec.processing_error = cad.DWG_MESSAGE if ext == ".dwg" else f"Type de fichier non pris en charge: {ext or mime}"
            db.commit()
            return {"file_id": file_rec.id, "status": "unsupported", "error": file_rec.processing_error}

        file_rec.page_count = 1
        file_rec.processing_status = "completed"
        db.add(ExtractedPage(file_id=file_rec.id, page_number=1, text=text,
                             classification="plan" if ext in cad.CAD_EXT else classify_page(text)))
        for chunk in chunk_text(text):
            db.add(DocumentChunk(file_id=file_rec.id, page_number=1, text=chunk))
        db.commit()
        return {"file_id": file_rec.id, "status": "completed", "chars": len(text)}
    except Exception as exc:
        file_rec.processing_status = "failed"
        file_rec.processing_error = str(exc)
        db.commit()
        return {"file_id": file_rec.id, "status": "failed", "error": str(exc)}
