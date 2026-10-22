"""Le « cerveau » de l'assistant : ce que le modèle reçoit (fil, historique), quand il réfléchit plus, comment il conclut."""
import os
import sys
from pathlib import Path
from types import SimpleNamespace as _NS

import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("UNIC_DATA_DIR", "/tmp/unic-test-data")
os.environ.setdefault("UNIC_NO_BACKGROUND", "1")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.main import app  # noqa: E402


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture
def seen(monkeypatch):
    """Remplace le modèle : garde chaque requête (messages + options), répond « ok »."""
    from app import orchestrator
    calls = []

    class R:
        text, provider, model, available, error, raw = "ok", "claude", "m", True, "", None

    def fake(msgs, **kw):
        calls.append((msgs, kw))
        return R()
    monkeypatch.setattr(orchestrator, "chat_complete", fake)
    monkeypatch.setattr(orchestrator, "provider_chain", lambda deep=False: [type("P", (), {"id": "claude"})()])
    return calls


def _turns(msgs):
    return [(m["role"], m["content"]) for m in msgs if m["role"] != "system"]


def test_la_demande_en_cours_part_une_seule_fois(client, seen):
    cid = client.post("/api/chat", json={"message": "Raconte une histoire courte"}).json()["conversation_id"]
    client.post("/api/chat", json={"message": "Une autre, plus drôle", "conversation_id": cid})
    turns = _turns(seen[-1][0])
    assert turns[-1] == ("user", "Une autre, plus drôle")
    assert [t for t in turns if t == ("user", "Une autre, plus drôle")] == [turns[-1]]   # avant : envoyée deux fois


def test_une_longue_conversation_garde_son_debut(client, seen):
    cid = client.post("/api/chat", json={"message": "Chantier de Mme Ndiaye à Ouakam : cloison de 12 m sur 2,50 m, deux faces"}).json()["conversation_id"]
    for i in range(7):
        client.post("/api/chat", json={"message": f"Question annexe sans rapport {chr(65 + i)}", "conversation_id": cid})
    client.post("/api/chat", json={"message": "Rappelle-moi la longueur de la cloison de Mme Ndiaye", "conversation_id": cid})
    msgs = seen[-1][0]
    variable = msgs[1]["content"]
    assert "PLUS TÔT DANS CETTE CONVERSATION" in variable and "cloison de 12 m sur 2,50 m" in variable
    assert len(_turns(msgs)) == 9                                   # 8 derniers échanges + la demande
    assert all("Mme Ndiaye à Ouakam" not in c for _, c in _turns(msgs))   # le début n'est PAS dans l'historique brut


def test_conversation_courte_sans_fil(client, seen):
    client.post("/api/chat", json={"message": "Bonjour, comment vas-tu ?"})
    assert "PLUS TÔT DANS CETTE CONVERSATION" not in seen[-1][0][1]["content"]


def test_les_sujets_a_enjeu_declenchent_la_reflexion_approfondie(client, seen):
    client.post("/api/chat", json={"message": "Prépare le devis de la cloison de 12 m pour M. Diop"})
    assert seen[-1][1]["deep"] is True
    client.post("/api/chat", json={"message": "Raconte-moi une blague"})
    assert seen[-1][1]["deep"] is False
    client.post("/api/chat", json={"message": "Réfléchis bien : une blague sur les chats"})
    assert seen[-1][1]["deep"] is True                              # demandée explicitement
    client.post("/api/chat", json={"message": "Combien coûte le devis ?", "voice": True})
    assert seen[-1][1]["deep"] is False                             # vocal : réponse rapide


def test_le_reglage_coupe_la_reflexion_automatique(client, seen):
    assert client.put("/api/settings", json={"ai_deep_stakes": False}).json()["ai_deep_stakes"] is False
    try:
        client.post("/api/chat", json={"message": "Prépare le devis de la cloison de 12 m"})
        assert seen[-1][1]["deep"] is False
        client.post("/api/chat", json={"message": "Réfléchis en profondeur à ce devis"})
        assert seen[-1][1]["deep"] is True                          # la demande explicite reste possible
    finally:
        client.put("/api/settings", json={"ai_deep_stakes": True})


# ---------- boucle d'outils ----------

def _resp(text="", stop="end_turn", content=None):
    return _NS(content=content or [_NS(type="text", text=text, citations=None)], stop_reason=stop, model="m",
               usage=_NS(input_tokens=10, output_tokens=5, cache_read_input_tokens=0, cache_creation_input_tokens=0,
                         server_tool_use=None))


def test_budget_d_outils_epuise_l_ia_conclut_au_lieu_de_rester_muette(monkeypatch):
    from app import ai
    from app.config import settings
    monkeypatch.setattr(settings, "anthropic_api_key", "secret-key")
    calls = []

    def create(**kw):
        calls.append(kw)
        if kw.get("tool_choice") == {"type": "none"}:
            return _resp("Fait : 8 recherches. Reste : le devis. Je continue ?")
        return _resp(stop="tool_use", content=[_NS(type="tool_use", id=f"tu_{len(calls)}", name="get_prices", input={})])
    fake = _NS(beta=_NS(messages=_NS(create=create)), messages=_NS(create=create))
    monkeypatch.setattr(ai.ClaudeAIProvider, "_client", lambda self: fake)
    ran = []
    r = ai.PROVIDERS["claude"].complete([{"role": "user", "content": "q"}], tools=[{"name": "get_prices"}],
                                        tool_handler=lambda name, args: ran.append(name) or {"ok": True})
    assert r.text.startswith("Fait : 8 recherches") and r.available
    assert len(ran) == ai.ClaudeAIProvider.MAX_ROUNDS                # l'appel en trop n'est PAS exécuté
    last = calls[-1]["messages"][-1]["content"][0]
    assert last["type"] == "tool_result" and last["is_error"] and "NON exécuté" in last["content"]


def test_la_reflexion_est_annoncee_pendant_le_flux():
    from app import ai

    class Stream:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def __iter__(self):
            yield _NS(type="content_block_start", content_block=_NS(type="thinking"))
            yield _NS(type="content_block_delta", delta=_NS(type="text_delta", text="Réponse"))

        def get_final_message(self):
            return _resp("Réponse")
    events = []
    client = _NS(beta=_NS(messages=_NS(stream=lambda **kw: Stream())))
    ai.ClaudeAIProvider._send_stream(client, {"model": "m", "messages": []}, events.append)
    assert {"t": "status", "text": "Je réfléchis…"} in events and {"t": "delta", "text": "Réponse"} in events


def test_un_oui_apres_une_proposition_de_devis_herite_de_l_enjeu(client, seen):
    cid = client.post("/api/chat", json={"message": "Prépare le devis de la cloison de 12 m pour M. Diop"}).json()["conversation_id"]
    client.post("/api/chat", json={"message": "oui vas-y", "conversation_id": cid})
    assert seen[-1][1]["deep"] is True                              # avant : « oui » = banalité, modèle le plus léger
    cid2 = client.post("/api/chat", json={"message": "Raconte une histoire de chat"}).json()["conversation_id"]
    client.post("/api/chat", json={"message": "oui encore", "conversation_id": cid2})
    assert seen[-1][1]["deep"] is False                             # une suite de blague reste légère


def test_vocal_la_confirmation_d_un_devis_ne_part_pas_sur_le_modele_leger(client, seen, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "anthropic_voice_model", "claude-haiku-4-5-20251001")
    cid = client.post("/api/chat", json={"message": "Fais le devis du plafond de M. Sow", "voice": True}).json()["conversation_id"]
    assert seen[-1][1].get("model") != settings.anthropic_voice_model
    client.post("/api/chat", json={"message": "oui", "conversation_id": cid, "voice": True})
    assert seen[-1][1].get("model") != settings.anthropic_voice_model   # le « oui » qui crée le devis : pas le modèle léger
    client.post("/api/chat", json={"message": "Quelle heure est-il à Paris ?", "voice": True})
    assert seen[-1][1].get("model") == settings.anthropic_voice_model    # simple conversation : le plus rapide


def test_la_voix_dit_les_montants_et_les_mesures_comme_un_humain():
    from app.voice import clean_for_speech
    assert clean_for_speech("Le devis fait **1 250 000 FCFA** pour 12 x 2,5 m.") == "Le devis fait 1250000 francs CFA pour 12 par 2,5 mètres."
    assert clean_for_speech("Il faut 24,5 m² de BA13 et 35 ml de bande.") == "Il faut 24,5 mètres carrés de B A 13 et 35 mètres linéaires de bande."
    assert clean_for_speech("RDV à 14h30, délai 48 h, 20 L et 17,5 kg.") == "RDV à 14 heures 30, délai 48 heures, 20 litres et 17,5 kilos."
    assert clean_for_speech("Mme Diop, malade, le 12/10.") == "Mme Diop, malade, le 12/10."


# ---------- universel : UniC ne se limite pas au placo ----------

@pytest.mark.parametrize("message", [
    "Envoie un mail à ma banque pour demander un rendez-vous",
    "Publie une annonce sur Facebook pour recruter un chauffeur",
    "Prépare un rapport de chantier sur la réunion de ce matin",
    "Statut du projet de livre que j'écris",
])
def test_les_demandes_universelles_vont_a_l_ia_pas_a_une_reponse_toute_faite(client, seen, message):
    before = len(seen)
    r = client.post("/api/chat", json={"message": message}).json()["message"]
    assert len(seen) == before + 1 and r["content"] == "ok"        # avant : texte figé, ou un PDF « rapport » vide


def test_l_enjeu_couvre_tous_les_domaines(client, seen):
    for msg, deep in (("Aide-moi à rédiger un contrat de travail pour un ouvrier", True),
                      ("Calcule mon impot sur 12 000 000 F de bénéfice", True),
                      ("J'ai des symptômes de paludisme, que faire ?", True),
                      ("Écris un poème sur Dakar", False)):
        client.post("/api/chat", json={"message": msg})
        assert seen[-1][1]["deep"] is deep, msg


def test_brouillon_de_nouvel_email(client):
    from app import agent
    from app.database import SessionLocal
    from app.models import EmailDraft
    with SessionLocal() as db:
        s = agent.AgentSession(db, None, {"owner_message": "écris un mail à ma banque"})
        out = s("save_email_draft", {"to": "conseiller@banque.sn", "subject": "Rendez-vous", "body": "Bonjour, je souhaite un rendez-vous."})
        d = db.get(EmailDraft, out["draft_id"])
        assert d.status == "draft" and d.to_addr == "conseiller@banque.sn" and s.cards[-1] == {"kind": "email", "id": d.id}
        out = s("save_email_draft", {"subject": "Bonjour", "body": "Texte"})                     # adresse inconnue : jamais inventée
        assert db.get(EmailDraft, out["draft_id"]).to_addr == "" and "compléter" in out["a"]
        assert "invalide" in s("save_email_draft", {"to": "pas-une-adresse", "subject": "x", "body": "y"})["error"]
        assert "introuvable" in s("save_email_draft", {"subject": "x", "body": "y", "document_number": "UC-0000"})["error"]


# ---------- hallucinations d'action ----------

def _scripted_chat(monkeypatch, replies):
    from app import orchestrator
    calls = []

    def fake(msgs, **kw):
        calls.append(msgs)
        text = replies[min(len(calls), len(replies)) - 1]

        class R:
            provider, model, available, error, raw = "claude", "m", True, "", None
        R.text = text
        return R()
    monkeypatch.setattr(orchestrator, "chat_complete", fake)
    monkeypatch.setattr(orchestrator, "provider_chain", lambda deep=False: [type("P", (), {"id": "claude"})()])
    return calls


def test_une_action_affirmee_sans_outil_est_corrigee(client, monkeypatch):
    calls = _scripted_chat(monkeypatch, ["J'ai créé le devis UC-2026-9999 pour M. Diop.",
                                         "Je n'ai pas encore créé le devis : donne-moi le lieu du chantier."])
    r = client.post("/api/chat", json={"message": "Fais le devis de M. Diop"}).json()["message"]
    assert len(calls) == 2 and "Contrôle automatique" in calls[1][-1]["content"]
    assert r["content"].startswith("Je n'ai pas encore créé le devis") and "⚠️" not in r["content"]


def test_si_l_ia_persiste_le_patron_est_averti(client, monkeypatch):
    _scripted_chat(monkeypatch, ["C'est noté : rendez-vous jeudi 10 h avec Awa."])
    r = client.post("/api/chat", json={"message": "Rendez-vous jeudi 10 h avec Awa"}).json()["message"]
    assert "⚠️ Vérification" in r["content"] and "le rendez-vous" in r["content"]


def test_une_reponse_sans_affirmation_n_est_pas_retouchee(client, monkeypatch):
    calls = _scripted_chat(monkeypatch, ["Il te faut 12 plaques. Veux-tu que je crée le devis ?"])
    client.post("/api/chat", json={"message": "Combien de plaques pour 5 x 2,5 m deux faces ?"})
    assert len(calls) == 1


# ---------- fichiers et pièces jointes ----------

def test_l_ia_cree_un_pdf_et_un_excel(client):
    from pathlib import Path
    from openpyxl import load_workbook
    from app import agent
    from app.database import SessionLocal
    from app.models import Artifact
    with SessionLocal() as db:
        s = agent.AgentSession(db, None, {})
        out = s("make_pdf", {"title": "Attestation de travaux", "content": "# Objet\nTravaux réalisés.\n- Plafond\n- **Cloisons**"})
        art = db.get(Artifact, out["artifact_id"])
        assert Path(art.path).read_bytes()[:5] == b"%PDF-" and art.filename == "Attestation_de_travaux.pdf" and s.files[-1]["id"] == art.id
        out = s("make_spreadsheet", {"title": "Budget", "sheets": [{"name": "Mois", "rows": [["Poste", "Montant"], ["Ciment", 50000], ["Total", "=SUM(B2:B2)"]]}]})
        ws = load_workbook(db.get(Artifact, out["artifact_id"]).path)["Mois"]
        assert ws["B2"].value == 50000 and ws["B3"].value == "=SUM(B2:B2)" and ws["A1"].font.bold
        assert "error" in s("make_spreadsheet", {"title": "Vide", "sheets": []})


def test_brouillon_avec_pieces_jointes_puis_envoi_apres_validation(client, monkeypatch):
    from app import agent, mailbox
    from app.database import SessionLocal
    sent = []
    monkeypatch.setattr(mailbox, "smtp_configured", lambda: True)
    monkeypatch.setattr(mailbox, "send", lambda to, subject, body, files=None: sent.append((to, subject, files)))
    with SessionLocal() as db:
        s = agent.AgentSession(db, None, {"owner_message": "envoie le rapport à la banque"})
        pdf = s("make_pdf", {"title": "Rapport", "content": "Texte du rapport."})["artifact_id"]
        out = s("save_email_draft", {"to": "conseiller@banque.sn", "subject": "Rapport", "body": "Bonjour, ci-joint.", "attachments": [pdf]})
        assert out["pieces_jointes"] == ["Rapport.pdf"]
        assert "introuvable" in s("save_email_draft", {"subject": "x", "body": "y", "attachments": ["pas-un-fichier"]})["error"]
    did = out["draft_id"]
    draft = next(d for d in client.get("/api/emails").json() if d["id"] == did)
    assert [a["filename"] for a in draft["attachments"]] == ["Rapport.pdf"]
    assert client.post(f"/api/mail/drafts/{did}/send").status_code == 409          # rien ne part sans l'approbation du patron
    assert client.post(f"/api/mail/drafts/{did}/approve").status_code == 200
    assert client.post(f"/api/mail/drafts/{did}/send").status_code == 200
    to, subject, files = sent[-1]
    assert to == "conseiller@banque.sn" and files[0][0] == "Rapport.pdf" and files[0][2][:5] == b"%PDF-"
