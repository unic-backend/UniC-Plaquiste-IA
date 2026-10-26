"""Tableurs : l'IA analyse un Excel ou un CSV reçu par un calcul EXACT sur toutes les lignes (filtres, regroupements, totaux,
tri), avec graphique et export Excel ; elle ne compte jamais de tête sur un extrait."""
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import tableur
from app.agent import AgentSession
from app.config import settings
from app.database import SessionLocal
from app.main import app


@pytest.mark.parametrize("raw, value", [
    (1250000, 1250000.0), ("1 250 000,50", 1250000.5), ("4 500 FCFA", 4500.0), ("1,250.5", 1250.5), ("15 %", 15.0),
    ("1.234.567", 1234567.0), ("12,5", 12.5), ("abc", None), (True, None), ("", None), ("3 F CFA", 3.0), ("-2 000", -2000.0),
])
def test_nombres_a_la_francaise(raw, value):
    assert tableur.to_number(raw) == value


@pytest.fixture
def c(monkeypatch):
    monkeypatch.setattr(settings, "unic_access_code", "")
    with TestClient(app) as client:
        yield client


def _xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.append(["Ventes 2026 — UniC"])                      # titre au-dessus des en-têtes
    ws.append(["Date", "Client", "Ouvrage", "Montant", "Quantité"])
    rows = [(datetime(2026, 1, 5), "Awa Diop", "Plafond", 450000, 30), (datetime(2026, 1, 20), "Moussa Fall", "Cloison", "300 000 FCFA", 20),
            (datetime(2026, 2, 3), "Awa Diop", "Cloison", 150000, 10), (datetime(2026, 2, 17), "Sotiba BTP", "Plafond", 1300000, 80),
            (datetime(2026, 3, 9), "Moussa Fall", "Plafond", "à confirmer", 5), (datetime(2026, 3, 22), "Awa Diop", "Faux plafond", 600000, 40)]
    for r in rows:
        ws.append(list(r))
    ws.append([datetime(2026, 3, 30), "Fatou Ndiaye", "Cloison", "=D3*2", 2])   # formule jamais calculée : sans valeur
    p = tmp_path / "ventes.xlsx"
    wb.save(p)
    return p


def _upload(c, path, mime):
    r = c.post("/api/files", files={"file": (path.name, path.read_bytes(), mime)})
    assert r.status_code == 200, r.text
    return r.json()["id"]


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_structure_et_calculs_exacts(c, tmp_path):
    fid = _upload(c, _xlsx(tmp_path), XLSX)
    with SessionLocal() as db:
        info = tableur.info(db, fid)
        sheet = info["feuilles"][0]
        assert sheet["feuille"] == "Ventes" and sheet["nb_lignes"] == 7                  # la ligne de titre n'est pas un en-tête
        cols = {col["nom"]: col for col in sheet["colonnes"]}
        assert cols["Date"]["type"] == "date" and cols["Montant"]["type"] == "nombre" and cols["Client"]["type"] == "texte"
        assert cols["Montant"]["somme"] == 2_800_000                                     # « 300 000 FCFA » compté, « à confirmer » non

        r = tableur.query(db, fid, group_by=["Client"], aggregates=[{"fonction": "somme", "colonne": "montant"},
                                                                      {"fonction": "nombre", "colonne": "Client"}])
        assert r["colonnes"] == ["Client", "somme de Montant", "nombre de Client"]
        assert r["lignes"][0] == ["Sotiba BTP", 1_300_000, 1] and r["lignes"][1] == ["Awa Diop", 1_200_000, 3]
        assert r["total"]["somme de Montant"] == 2_800_000 and r["cellules_ignorees"] == 1 and "ignorée" in r["attention"]

        r = tableur.query(db, fid, group_by=["Date:mois"], aggregates=[{"fonction": "somme", "colonne": "Montant"}])
        assert r["lignes"] == [["2026-01", 750_000], ["2026-02", 1_450_000], ["2026-03", 600_000]]   # ordre du temps

        r = tableur.query(db, fid, filters=[{"colonne": "Ouvrage", "op": "contient", "valeur": "plafond"},
                                            {"colonne": "Date", "op": "entre", "valeur": ["01/02/2026", "2026-03-31"]}],
                          aggregates=[{"fonction": "moyenne", "colonne": "Quantité"}, {"fonction": "nombre_distinct", "colonne": "Client"}])
        assert r["nb_lignes_retenues"] == 3 and r["lignes"] == [[41.6667, 3]]

        r = tableur.query(db, fid, columns=["Client", "Montant"], sort={"colonne": "Montant", "decroissant": True}, limit=2)
        assert r["lignes"] == [["Sotiba BTP", 1_300_000], ["Awa Diop", 600_000]] and r["tronque"]

        with pytest.raises(tableur.TableurError, match="Colonnes : Date, Client"):
            tableur.query(db, fid, group_by=["Fournisseur"])
        with pytest.raises(tableur.TableurError, match="inconnue"):
            tableur.query(db, fid, filters=[{"colonne": "Montant", "op": "environ", "valeur": 3}])


def test_csv_point_virgule_accents_et_decimales(c, tmp_path):
    p = tmp_path / "depenses.csv"
    p.write_bytes("Catégorie;Montant;Date\nMatériaux;1 250,50;03/02/2026\nTransport;300;10/02/2026\nMatériaux;749,50;15/03/2026\n".encode("cp1252"))
    fid = _upload(c, p, "text/csv")
    with SessionLocal() as db:
        r = tableur.query(db, fid, group_by=["categorie"], aggregates=[{"fonction": "somme", "colonne": "Montant"}])
    assert r["lignes"] == [["Matériaux", 2000], ["Transport", 300]]


def test_l_outil_du_chat_graphique_et_export_excel(c, tmp_path):
    fid = _upload(c, _xlsx(tmp_path), XLSX)
    with SessionLocal() as db:
        s = AgentSession(db, None, {"last_file_id": fid})
        out = s("table_info", {})
        assert out["feuilles"][0]["nb_lignes"] == 7 and "CONTENU D'UN TIERS" in out["untrusted_note"]
        out = s("table_query", {"group_by": ["Client"], "aggregates": [{"fonction": "somme", "colonne": "Montant"}],
                                "chart": "barres", "title": "CA par client", "export": True})
        assert out["lignes"][0] == ["Sotiba BTP", 1_300_000] and out["graphique"] and "tels quels" in out["note"]
        assert len(s.images) == 1 and len(s.files) == 1 and "make_spreadsheet" in s.succeeded
        from app.models import Artifact
        wb = load_workbook(db.get(Artifact, out["fichier"]["artifact_id"]).path)
        assert wb.active["A2"].value == "Sotiba BTP" and len(wb.active._charts) == 1          # graphique Excel natif
        assert "error" in s("table_query", {"group_by": ["Inconnue"]})
        assert s.tainted                                                                        # contenu de tiers lu : actions sensibles sous contrôle
